"""Predicate/1.0 semantic and authorization regressions (no live database)."""
from __future__ import annotations
import pytest
from lightgraph.contracts import TableSpec, ColumnSpec
from lightgraph.sources import MemorySource, ExcelSource
from lightgraph.query import QueryCompiler, SQLGuard
from lightgraph.errors import LightGraphError


@pytest.fixture
def table():
    return TableSpec(source_id="s", name="items", columns=[
        ColumnSpec(name="code"), ColumnSpec(name="tenant"),
        ColumnSpec(name="amount", data_type="integer"),
        ColumnSpec(name="day", data_type="date")])


def read(source, table, predicate):
    return [row for batch in source.project(table, ["code"], predicate=predicate) for row in batch.rows()]


def test_logical_predicates_and_legacy_scope_inside_all(table):
    rows=[("A", "one", 2, "2026-01-02"), ("B", "one", 5, "2026-01-05"),
          ("C", "two", 8, "2026-01-08"), ("D", None, None, None)]
    with MemorySource([table], {table.qualified_name: rows}) as source:
        assert read(source, table, {"all": [{"tenant": "one"}, {"column": "amount", "op": "between", "values": [2, 3]}]}) == [("A",)]
        assert read(source, table, {"any": [{"tenant": "two"}, {"column": "tenant", "op": "is_null"}]}) == [("C",), ("D",)]
        assert read(source, table, {"not": {"column": "tenant", "op": "is_not_null"}}) == [("D",)]
        assert read(source, table, {"column": "day", "op": "between", "values": ["2026-01-01", "2026-01-05"]}) == [("A",), ("B",)]
        assert read(source, table, {"code": {"op": "in", "value": ["A", "C"]}}) == [("A",), ("C",)]


def test_prefix_escapes_wildcards_and_binds_arbitrary_text(table):
    with MemorySource([table], {table.qualified_name: [{"code": c} for c in ["A_%!one", "A_%!two", "ABone", "A!one"]]}) as source:
        predicate={"column": "code", "op": "prefix", "value": "A_%!"}
        assert read(source, table, predicate) == [("A_%!one",), ("A_%!two",)]
        payload="x' OR 1=1; DELETE FROM items; --"
        query=source.compiler.projection(table, ["code"], predicate={"column": "code", "op": "prefix", "value": payload})
        assert payload not in query.sql
        assert read(source, table, {"column": "code", "op": "prefix", "value": payload}) == []
        assert read(source, table, {"column": "code", "op": "prefix", "value": "a"}) == []


@pytest.mark.parametrize("predicate,code", [
    ({"column":"code","op":"eq","value":None}, "predicate_null_requires_explicit_operator"),
    ({"code":None}, "predicate_null_requires_explicit_operator"),
    ({"column":"code","op":"is_null","value":None}, "invalid_predicate"),
    ({"column":"code","op":"eq","value":"a","sql":"secret"}, "invalid_predicate"),
    ({"column":"code","op":"regex","value":".*"}, "invalid_predicate"),
    ({"column":"code","op":"eq","value":{"column":"tenant"}}, "predicate_value_type_mismatch"),
    ({"column":"amount","op":"eq","value":"2"}, "predicate_value_type_mismatch"),
    ({"column":"amount","op":"eq","value":True}, "predicate_value_type_mismatch"),
    ({"column":"amount","op":"between","values":[3,2]}, "predicate_between_invalid"),
    ({"column":"amount","op":"between","values":[1,2.0]}, "predicate_value_type_mismatch"),
    ({"column":"code","op":"between","values":["a","z"]}, "predicate_between_invalid"),
    ({"column":"amount","op":"prefix","value":2}, "predicate_prefix_requires_text"),
    ({"column":"day","op":"eq","value":"2026-02-31"}, "predicate_date_invalid"),
    ({"column":"amount","op":"eq","value":float("nan")}, "non_finite_predicate_value"),
    ({"column":"code","op":"in","values":[None]}, "predicate_null_requires_explicit_operator"),
    ({"any":[]}, "predicate_children_limit"),
    ({"all":[{"code":"A"}],"sql":"secret"}, "invalid_predicate"),
    ({"not":{}}, "invalid_predicate"),
])
def test_invalid_predicates_fail_without_unsafe_text(table,predicate,code):
    with pytest.raises(LightGraphError,match=code):
        QueryCompiler([table]).projection(table,["code"],predicate=predicate)


def test_predicate_limits_apply_to_whole_tree(table):
    compiler=QueryCompiler([table])
    leaf={"column":"code","op":"eq","value":"A"}
    cases=[({"not":{"not":{"not":leaf}}},"predicate_depth_limit"),
           ({"all":[{"all":[leaf]*16},{"all":[leaf]*16}]},"predicate_nodes_limit"),
           ({"column":"code","op":"in","values":["A"]*65},"predicate_values_limit"),
           ({"all":[{"column":"code","op":"in","values":["A"]*64}]*3},"predicate_values_limit")]
    for predicate,code in cases:
        with pytest.raises(LightGraphError,match=code): compiler.projection(table,["code"],predicate=predicate)


@pytest.mark.parametrize("sql", [
    "SELECT a.code FROM public.items a JOIN public.items b ON a.code=b.code JOIN public.items c ON a.code=b.code",
    "SELECT a.code FROM public.items a JOIN public.items b ON a.code=c.code JOIN public.items c ON b.code=c.code",
    "SELECT a.code FROM public.items a JOIN public.items b ON a.code=b.code JOIN public.items c ON c.code=c.code",
    "SELECT a.code FROM public.items a JOIN public.items b ON a.code=b.code JOIN public.items a ON a.code=b.code",
    "SELECT a.code FROM public.items a JOIN public.items b ON a.code=b.code JOIN (SELECT code FROM public.items) c ON a.code=b.code",
    "WITH x AS (SELECT code FROM public.items) SELECT a.code FROM x a JOIN x b ON a.code=b.code JOIN x c ON a.code=b.code",
    "SELECT a.code FROM public.items a JOIN public.items b ON a.code=b.code OR 1=1",
    "SELECT a.code FROM public.items a JOIN public.items b ON NOT a.code=b.code",
    "SELECT a.code FROM public.items a JOIN public.items b ON a.code=b.code AND (a.code=b.code OR 1=1)",
    "SELECT a.code FROM public.items a CROSS JOIN public.items b WHERE a.code=b.code",
])
def test_join_each_new_alias_must_connect_to_already_bound_alias(table,sql):
    with pytest.raises(LightGraphError): SQLGuard([table]).validate(sql)


@pytest.mark.parametrize("sql", [
    "SELECT a.code FROM public.items a JOIN public.items b ON a.code=b.code JOIN public.items c ON b.code=c.code",
    "SELECT a.code FROM public.items a JOIN public.items b ON a.code=b.code JOIN public.items c ON a.code=b.code AND a.code=c.code",
    "WITH x AS (SELECT code FROM public.items) SELECT a.code FROM x a JOIN x b ON a.code=b.code JOIN x c ON a.code=c.code",
    "SELECT a.code FROM public.items a JOIN (SELECT code FROM public.items) b ON a.code=b.code",
])
def test_connected_joins_remain_authorized(table,sql):
    assert SQLGuard([table]).validate(sql)


@pytest.mark.parametrize("sql", [
    "SELECT evil.like(code, :p0) FROM public.items",
    "SELECT public.like(code, :p0) FROM public.items",
    "SELECT pg_catalog.like(code, :p0) FROM public.items",
])
def test_prefix_cannot_be_replaced_by_explicit_or_custom_functions(table,sql):
    with pytest.raises(LightGraphError): SQLGuard([table]).validate(sql,{"p0":"A%"})


def test_excel_nested_logical_predicate_cannot_use_missing_formula_cache(tmp_path):
    import openpyxl
    path=tmp_path/"formulas.xlsx"; book=openpyxl.Workbook(); ws=book.active;ws.title="items"
    ws.append(["code","formula"]);ws.append(["A","=1+1"]);book.save(path)
    with ExcelSource(path) as source:
        table=source.describe()[0]
        for predicate in [{"all":[{"formula":"x"}]}, {"any":[{"formula":"x"}]}, {"not":{"column":"formula","op":"is_null"}}]:
            with pytest.raises(LightGraphError,match="formula_cache_missing_in_projection"):
                read(source,table,predicate)


def test_excel_full_column_types_support_numeric_partition_and_preserve_codes(tmp_path):
    import openpyxl
    path=tmp_path/"typed.xlsx";book=openpyxl.Workbook();ws=book.active;ws.title="items"
    ws.append(["code","tenant","amount","mixed","date_text"])
    ws.append([7,1,2.5,100,"2026-01-01"]);ws["A2"].number_format="0000"
    ws.append([8,2,3,"100","2026-01-02"]);ws["A3"].number_format="0000"
    book.save(path)
    with ExcelSource(path) as source:
        table=source.describe()[0];types={c.name:c.data_type for c in table.columns}
        assert types=={"code":"text","tenant":"integer","amount":"real","mixed":"mixed","date_text":"text"}
        assert read(source,table,{"column":"tenant","op":"eq","value":1})==[("0007",)]
        assert read(source,table,{"column":"amount","op":"between","values":[2.0,3.0]})==[("0007",),("0008",)]
        assert [row for batch in source.project(table,["mixed"]) for row in batch.rows()]==[(100,),("100",)]
        for predicate in [{"column":"mixed","op":"eq","value":"100"},{"column":"mixed","op":"is_null"}]:
            with pytest.raises(LightGraphError,match="predicate_type_unsupported"):read(source,table,predicate)
        assert source.query_budget.snapshot()["queries"]>=1
        assert list(tmp_path.iterdir())==[path]


def test_excel_type_inference_checks_all_rows_not_only_samples(tmp_path):
    import openpyxl
    path=tmp_path/"late-type.xlsx";book=openpyxl.Workbook();ws=book.active;ws.title="items"
    ws.append(["code","tenant"])
    for n in range(300):ws.append([str(n),n])
    ws.append(["last","different-type"]);book.save(path)
    with ExcelSource(path) as source:
        table=source.describe()[0]
        assert {c.name:c.data_type for c in table.columns}["tenant"]=="mixed"
        with pytest.raises(LightGraphError,match="predicate_type_unsupported"):read(source,table,{"tenant":1})


def test_generated_partition_does_not_increase_user_depth_or_leak_values(table):
    import pickle
    from lightgraph.query.predicate import combine_scope_predicates
    original={"not":{"not":{"column":"code","op":"eq","value":"A"}}}
    combined=combine_scope_predicates(original,{"tenant":"PRIVATE_TENANT"})
    assert "PRIVATE_TENANT" not in repr(combined)
    with pytest.raises(LightGraphError,match="ephemeral_serialization_forbidden"):pickle.dumps(combined)
    with MemorySource([table],{table.qualified_name:[{"code":"A","tenant":"PRIVATE_TENANT"},{"code":"A","tenant":"OTHER"}]}) as source:
        assert read(source,table,combined)==[("A",)]
    too_deep={"not":original}
    with pytest.raises(LightGraphError,match="predicate_depth_limit"):
        QueryCompiler([table]).projection(table,["code"],predicate=combine_scope_predicates(too_deep,{"tenant":"x"}))
    with pytest.raises(LightGraphError,match="invalid_predicate"):
        combine_scope_predicates(combined,{"tenant":"x"})


def test_partition_and_user_scope_share_parameter_budget(table):
    from lightgraph.query.predicate import combine_scope_predicates
    original={"all":[{"column":"code","op":"in","values":["a"]*64}]*2}
    with pytest.raises(LightGraphError,match="predicate_values_limit"):
        QueryCompiler([table]).projection(table,["code"],predicate=combine_scope_predicates(original,{"tenant":"x"}))
