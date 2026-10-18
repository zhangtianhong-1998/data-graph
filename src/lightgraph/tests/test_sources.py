from __future__ import annotations
import pickle
import sqlite3
from pathlib import Path
import pytest
from lightgraph.contracts import TableSpec, ColumnSpec, QueryLimits
from lightgraph.sources import MemorySource, ExcelSource, EphemeralBatch
from lightgraph.query import SQLGuard
from lightgraph.errors import LightGraphError

@pytest.fixture
def table():
    return TableSpec(source_id="s", name="items", columns=[ColumnSpec(name=c) for c in ["tenant", "code", "name", "parent", "secret"]])

@pytest.fixture
def source(table):
    with MemorySource([table], {"public.items": [
        {"tenant":"A", "code":"001", "name":"Sales", "parent":"ROOT", "secret":"UNPUBLISHED_MARKER"},
        {"tenant":"A", "code":"001", "name":"Sales", "parent":"ROOT"},
        {"tenant":"B", "code":"001", "name":"Expenses", "parent":"ROOT"},
        {"tenant":"B", "code":"002", "name":None, "parent":None},
    ]}) as s: yield s

def read(source, table, columns, **kwargs):
    result=[]
    for batch in source.project(table, columns, **kwargs): result.extend(batch.rows())
    return result

def test_joint_projection_preserves_pairs_and_scopes(source, table):
    assert set(read(source, table, ["tenant","code","name"])) == {("A","001","Sales"),("B","001","Expenses"),("B","002",None)}
    assert read(source,table,["code","name"],predicate={"tenant":"A"}) == [("001","Sales")]

def test_predicate_values_never_become_sql(source, table):
    text="x' OR 1=1; DELETE FROM items; --"
    q=source.compiler.projection(table,["code"],predicate={"tenant":text})
    assert text not in q.sql
    assert read(source,table,["code"],predicate={"tenant":text}) == []
    assert text not in repr(q)
    with pytest.raises(LightGraphError,match="serialization"): pickle.dumps(q)

def test_limit_detects_one_beyond_and_exact_limit_is_complete(source, table):
    assert len(read(source,table,["tenant","code"],limit=3)) == 3
    with pytest.raises(LightGraphError,match="result_truncated"): read(source,table,["tenant","code"],limit=2)

def test_batch_released_and_no_pickle():
    b=EphemeralBatch(["a"],[("do-not-print",)])
    assert "do-not-print" not in repr(b)
    with pytest.raises(LightGraphError,match="serialization"): pickle.dumps(b)
    b.release()
    with pytest.raises(LightGraphError,match="expired"): list(b.rows())

def test_batch_byte_budget():
    with pytest.raises(LightGraphError,match="memory_budget"): EphemeralBatch(["a"],[("x"*100,)],max_bytes=20)

def test_fd_scope_and_null_are_distinct(source, table):
    assert source.check_fd(table,["code"],["name"])["violating_groups"] == 1
    check=source.check_fd(table,["tenant","code"],["name"])
    assert check["violating_groups"] == 0
    assert check["all_null_groups"] == 1

def test_reference_duplicates_not_ambiguity(table):
    refs=TableSpec(source_id="s",name="refs",columns=[ColumnSpec(name=c) for c in ["tenant","ref"]])
    with MemorySource([table,refs],{"public.items":[{"tenant":"A","code":"1","name":"X"}]*20 + [{"tenant":"B","code":"1","name":"Y"},{"tenant":"B","code":"1","name":"Z"}],"public.refs":[{"tenant":"A","ref":"1"},{"tenant":"B","ref":"1"},{"tenant":"A","ref":"2"}]}) as s:
        result=s.check_reference(refs,["tenant","ref"],table,["tenant","code"],discriminator_columns=["name"])
        assert result == {"source_ref_tuples":3,"null_refs":0,"unmatched":1,"ambiguous":1,"unique_matches":1}

@pytest.mark.parametrize("sql",[
    "SELECT code FROM public.items; DELETE FROM public.items",
    "WITH x AS (DELETE FROM public.items RETURNING code) SELECT code FROM x",
    "SELECT code INTO x FROM public.items",
    "COPY public.items TO '/tmp/test'",
    "SELECT pg_read_file('/etc/passwd')",
    "SELECT nextval('seq')",
    "SELECT code FROM public.items FOR UPDATE",
    "SELECT code FROM unauthorized.items",
    "SELECT excluded FROM public.items",
    "SELECT secret FROM public.items",
    "SELECT * FROM public.items",
    "SELECT code FROM public.items OFFSET 1000000",
    "WITH RECURSIVE x AS (SELECT 1 UNION ALL SELECT 2) SELECT * FROM x",
    "EXPLAIN ANALYZE SELECT code FROM public.items",
    "SELECT code FROM public.items WHERE malicious_udf(code)",
    "SELECT code FROM public.items; SET role admin",
    "SELECT a.code FROM public.items a JOIN public.items b ON 1=1",
    "SELECT a.code FROM public.items a JOIN public.items b ON a.code=a.code",
])
def test_sql_guard_denies_unsafe_and_unauthorized(sql,table):
    authorized=table.model_copy(deep=True); authorized.columns=[c for c in authorized.columns if c.name!="secret"]
    with pytest.raises(LightGraphError): SQLGuard([authorized]).validate(sql)

def test_guard_rechecks_cross_source_authority(table):
    other=TableSpec(source_id="s",name="other",columns=[ColumnSpec(name="code")])
    with MemorySource([table],{"public.items":[]}) as a, MemorySource([other],{"public.other":[]}) as b:
        q=a.compiler.projection(table,["code"])
        with pytest.raises(LightGraphError,match="unauthorized_object"): list(b.execute(q))

def test_excel_does_not_execute_formulas_and_preserves_leading_zero(tmp_path):
    import openpyxl
    p=tmp_path/"input.xlsx"; book=openpyxl.Workbook(); ws=book.active; ws.title="items"
    ws.append(["code","name","formula"]); ws.append([7,"Sales","=HYPERLINK(\"https://bad.invalid\",\"x\")"]); ws["A2"].number_format="0000"; book.save(p)
    with ExcelSource(p,source_id="s") as s:
        t=s.describe()[0]
        assert t.qualified_name == "public.items"
        assert read(s,t,["code","name"]) == [("0007","Sales")]
        with pytest.raises(LightGraphError,match="formula_cache_missing_in_projection"):
            read(s,t,t.column_names)
        with pytest.raises(LightGraphError,match="formula_cache_missing_in_projection"):
            read(s,t,["code"],predicate={"formula":None})
        assert s.issues[0]["code"] == "formula_cache_missing"
        assert set(tmp_path.iterdir()) == {p}

def test_excel_hidden_rows_and_columns_are_explicit(tmp_path):
    import openpyxl
    p=tmp_path/"input.xlsx"; b=openpyxl.Workbook(); w=b.active; w.title="t"
    w.append(["code","hidden"]); w.append(["A","x"]);w.append(["B","y"])
    w.column_dimensions["B"].hidden=True; w.row_dimensions[3].hidden=True;b.save(p)
    with ExcelSource(p) as s:
        t=s.describe()[0]; assert t.column_names==["code"]
        assert read(s,t,["code"])==[("A",)]

def test_source_errors_and_stats_do_not_contain_payload(table):
    with MemorySource([table],{"public.items":[{"code":"PRIVATE_MARKER"}]},limits=QueryLimits(max_batch_bytes=4)) as s:
        with pytest.raises(LightGraphError) as caught: read(s,table,["code"])
        assert "PRIVATE_MARKER" not in str(caught.value)
        assert "PRIVATE_MARKER" not in str(s.stats.safe_dict())

def test_all_probe_shapes(source,table):
    assert source.probe("profile_column",table,{"column":"code"})["rows"]==4
    assert source.probe("profile_column_group",table,{"columns":["tenant","code"]})["distinct_tuples"]==3
    assert source.probe("check_value_domain",table,{"column":"code","allowed":["001"]})["out_of_domain"]==1
    assert source.probe("path_probe",table,{"columns":["parent","code"]})["rows"]==4
    assert len(source.probe("level_probe",table,{"columns":["parent","code"]})["per_level"])==2
    assert source.probe("compare_attribute_role",table,{"column":"code","target":"public.items","target_column":"code"})["overlap_values"]==2
    batches=source.probe("unpivot_probe",table,{"levels":[{"code":"parent"},{"code":"code","label":"name"}]})
    assert sum(len(list(batch.rows())) for batch in batches)==4
    batches=source.probe("sample_strata",table,{"columns":["tenant","code"],"limit":2})
    assert sum(len(list(batch.rows())) for batch in batches)==2

def test_query_budget_is_shared_across_sources(table):
    from lightgraph.sources import QueryBudget
    budget=QueryBudget(max_queries=2)
    with MemorySource([table],{"public.items":[{"code":"A"}]},query_budget=budget) as a, MemorySource([table],{"public.items":[{"code":"B"}]},query_budget=budget) as b:
        assert read(a,table,["code"])==[("A",)]
        assert read(b,table,["code"])==[("B",)]
        with pytest.raises(LightGraphError,match="query_budget_exhausted"): read(a,table,["code"])
        assert budget.snapshot()["queries"]==2

def test_global_return_bytes_and_deadline_are_enforced(table):
    from lightgraph.sources import QueryBudget
    import time
    budget=QueryBudget(max_returned_bytes=18)
    with MemorySource([table],{"public.items":[{"code":"ABCDE"}]},query_budget=budget) as s:
        read(s,table,["code"])
        with pytest.raises(LightGraphError,match="returned_bytes_budget_exhausted"): read(s,table,["code"])
    budget=QueryBudget()
    with MemorySource([table],{"public.items":[]},query_budget=budget) as s:
        budget.deadline=time.monotonic()-1
        with pytest.raises(LightGraphError,match="run_time_budget_exhausted"): read(s,table,["code"])

def test_install_budget_imports_prior_consumption(table):
    from lightgraph.sources import QueryBudget
    with MemorySource([table],{"public.items":[]}) as s:
        read(s,table,["code"])
        budget=QueryBudget(max_queries=1)
        s.install_budget(budget)
        with pytest.raises(LightGraphError,match="query_budget_exhausted"): read(s,table,["code"])

def test_install_shared_prior_budget_once_and_keep_earlier_deadline(table):
    from lightgraph.sources import QueryBudget
    original=QueryBudget(max_run_seconds=60)
    with MemorySource([table],{"public.items":[{"code":"A"}]},query_budget=original) as a, MemorySource([table],{"public.items":[{"code":"B"}]},query_budget=original) as b:
        read(a,table,["code"]);read(b,table,["code"])
        replacement=QueryBudget(max_queries=3,max_run_seconds=3600)
        a.install_budget(replacement);b.install_budget(replacement)
        assert replacement.queries==2
        assert replacement.returned_rows==2
        assert replacement.deadline==original.deadline
        read(a,table,["code"])
        with pytest.raises(LightGraphError,match="query_budget_exhausted"):read(b,table,["code"])

def test_excel_formula_pending_columns_survive_multiple_sheets_and_capped_issues(tmp_path):
    import openpyxl
    path=tmp_path/"formula.xlsx";book=openpyxl.Workbook();a=book.active;a.title="A"
    a.append(["code","l1","l2"])
    for i in range(35):a.append([str(i),"root","=1+1"])
    b=book.create_sheet("B");b.append(["code","other"]);b.append(["X","=1+1"]);book.save(path)
    with ExcelSource(path) as source:
        tables={t.name:t for t in source.describe()}
        assert len(source.issues)==32
        assert source.issue_counts["formula_cache_missing"]==36
        for name,column in (("A","l2"),("B","other")):
            with pytest.raises(LightGraphError,match="formula_cache_missing_in_projection"):
                read(source,tables[name],["code",column])
        assert len(read(source,tables["A"],["code","l1"]))==35

@pytest.mark.parametrize('value',[float('nan'),float('inf'),float('-inf')])
def test_non_finite_memory_values_are_rejected_before_sqlite_coerces_nan(table,value):
    with pytest.raises(LightGraphError,match='non_finite_source_value'):
        MemorySource([table],{'public.items':[{'code':value}]})

@pytest.mark.parametrize('value',[b'PRIVATE_BINARY',memoryview(b'PRIVATE_BINARY')])
def test_binary_driver_values_have_explicit_safe_error(value):
    from lightgraph.sources.base import cursor_batches,QueryStats
    class Cursor:
        def fetchone(self):return (value,)
    with pytest.raises(LightGraphError) as caught:
        list(cursor_batches(Cursor(),limits=QueryLimits(),columns=('code',),row_limit=1,stats=QueryStats()))
    assert str(caught.value)=='source_value_type_unsupported'
    assert 'PRIVATE_BINARY' not in str(caught.value)

@pytest.mark.parametrize('text',['NaN','Infinity','-Infinity'])
def test_non_finite_decimal_driver_values_are_rejected(text):
    from decimal import Decimal
    from lightgraph.sources.base import cursor_batches,QueryStats
    class Cursor:
        def fetchone(self):return (Decimal(text),)
    with pytest.raises(LightGraphError,match='non_finite_source_value'):
        list(cursor_batches(Cursor(),limits=QueryLimits(),columns=('code',),row_limit=1,stats=QueryStats()))

@pytest.mark.parametrize('operator,args',[
    ('profile_column',{'column':'evil'}),
    ('profile_column_group',{'columns':['evil']}),
    ('check_fd',{'keys':['evil'],'values':['name']}),
    ('check_value_domain',{'column':'evil','allowed':['x']}),
    ('sample_strata',{'columns':['evil']}),
    ('check_reference',{'source_columns':['evil'],'target':'public.items','target_columns':['code']}),
    ('distinct_projection',{'columns':['evil']}),
    ('path_probe',{'columns':['code','evil']}),
    ('level_probe',{'columns':['code','evil']}),
    ('unpivot_probe',{'levels':[{'code':'evil'}]}),
    ('compare_attribute_role',{'column':'evil','target':'public.items','target_column':'code'}),
])
def test_all_registered_operators_reject_identifier_injection(source,table,operator,args):
    def replace(v):
        if isinstance(v,str) and v=='evil':return 'code); DELETE FROM public.items; --'
        if isinstance(v,list):return [replace(x) for x in v]
        if isinstance(v,dict):return {k:replace(x) for k,x in v.items()}
        return v
    with pytest.raises(LightGraphError):
        result=source.probe(operator,table,replace(args))
        if not isinstance(result,dict):list(result)
    assert source.profile(table,['code'])['code']['rows']==4


def test_role_and_level_probes_respect_scope_predicate(source,table):
    result=source.probe('level_probe',table,{'columns':['parent','code']},predicate={'tenant':'A'})
    assert result['per_level']['code']['rows']==2
    result=source.probe('compare_attribute_role',table,{'column':'code','target':'public.items','target_column':'code'},predicate={'tenant':'A'})
    assert result['overlap_values']==1


def test_internal_aggregate_names_do_not_collide_with_source_columns():
    t=TableSpec(source_id='s',name='collision',columns=[ColumnSpec(name=c) for c in ['c','n','freq','missing','_lg_bindings']])
    rows=[{'c':'KEY','n':'A','freq':'text','missing':'text','_lg_bindings':'KEY'}, {'c':'KEY','n':'B','freq':'text','missing':'text','_lg_bindings':'KEY'}]
    with MemorySource([t],{'public.collision':rows}) as s:
        fd=s.check_fd(t,['c'],['n']);assert fd['violating_groups']==1 and fd['excess_values']==1
        group=s.probe('profile_column_group',t,{'columns':['freq','missing']})
        assert group['duplicate_groups']==1 and group['null_key_rows']==0
        ref=s.check_reference(t,['c'],t,['_lg_bindings'],discriminator_columns=['n'])
        assert ref['ambiguous']==1

def test_unreferenced_custom_type_does_not_block_authorized_safe_projection():
    t=TableSpec(source_id='s',name='typed',columns=[ColumnSpec(name='code'),ColumnSpec(name='private_object',data_type='USER-DEFINED')])
    guard=SQLGuard([t])
    guard.validate('SELECT code FROM public.typed')
    guard.validate('SELECT COUNT(*) FROM public.typed')
    with pytest.raises(LightGraphError,match='untrusted_column_type'):
        guard.validate('SELECT private_object FROM public.typed')
