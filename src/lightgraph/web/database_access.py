"""Bounded database onboarding before a table allowlist exists.

Only fixed, parameterized catalog queries run here. Agent-generated SQL still
uses the extraction guard; onboarding never exposes an arbitrary SQL endpoint
or reads instance rows. Connection secrets exist only in request/session memory.
"""
from __future__ import annotations

from contextlib import contextmanager
from io import BytesIO
import ipaddress
import re
import sys
from typing import Literal
from urllib.parse import parse_qsl, quote, unquote, urlsplit
from zipfile import ZipFile

from pydantic import BaseModel, ConfigDict, Field, SecretStr, field_validator, model_validator

from lightgraph.contracts import QueryLimits
from lightgraph.errors import LightGraphError
from lightgraph.sources.mysql import _connect_options

TABLE_LIMIT = 1000
CONNECTION_TIMEOUT_SECONDS = 5
MAX_EXCEL_BYTES = 5 * 1024 * 1024
MAX_EXCEL_EXPANDED_BYTES = 20 * 1024 * 1024
_IDENTIFIER = re.compile(r"[^\x00-\x1f\x7f./\\]{1,128}\Z", re.UNICODE)
_SYSTEM_SCHEMAS = {"mysql", "sys", "information_schema", "performance_schema", "pg_catalog", "pg_toast"}
_GAUSS_JDBC_PREFIXES = ("jdbc:gaussdb://", "jdbc:postgresql://")
GaussSSLMode = Literal["disable", "allow", "prefer", "require", "verify-ca", "verify-full"]
_GAUSS_SSL_MODES = {"disable", "allow", "prefer", "require", "verify-ca", "verify-full"}


class DatabaseAccessError(LightGraphError):
    """Safe codes only; driver diagnostics can contain credentials or SQL."""


class ConnectionSpec(BaseModel):
    model_config = ConfigDict(extra="forbid", hide_input_in_errors=True)

    kind: Literal["gauss_database", "mysql_database"]
    dsn: SecretStr | None = Field(default=None, repr=False)
    host: str | None = Field(default=None, max_length=253)
    port: int | None = Field(default=None, ge=1, le=65535)
    database: str | None = Field(default=None, max_length=128)
    username: str | None = Field(default=None, max_length=128)
    password: SecretStr | None = Field(default=None, repr=False)
    schema_name: str | None = Field(default=None, max_length=128, alias="schema")
    sslmode: GaussSSLMode | None = None

    @field_validator("kind", mode="before")
    @classmethod
    def normalize_kind(cls, value):
        return {"gauss": "gauss_database", "mysql": "mysql_database"}.get(value, value) if isinstance(value, str) else value

    @model_validator(mode="after")
    def one_connection_form(self):
        if self.kind != "gauss_database" and self.sslmode is not None:
            raise ValueError("database_sslmode_not_supported")
        if self.dsn is not None:
            if not self.dsn.get_secret_value() or len(self.dsn.get_secret_value()) > 8192:
                raise ValueError("database_connection_invalid")
            jdbc = self.kind == "gauss_database" and self.dsn.get_secret_value().strip().startswith(_GAUSS_JDBC_PREFIXES)
            if any(value is not None for value in (self.host, self.port, self.database)) or (
                not jdbc and any(value is not None for value in (self.username, self.password))
            ):
                raise ValueError("database_connection_forms_conflict")
        elif not self.host or not self.database or not self.username:
            raise ValueError("database_connection_fields_required")
        if self.password is not None and len(self.password.get_secret_value()) > 4096:
            raise ValueError("database_connection_invalid")
        return self


def _identifier(value: str) -> str:
    # Names are never SQL text: catalog filters bind parameters and downstream
    # compilers quote identifiers. Dots remain unsupported by schema.table IDs.
    if not isinstance(value, str) or not value.strip() or not _IDENTIFIER.fullmatch(value):
        raise DatabaseAccessError("database_identifier_invalid")
    return value


def _schema(value: str) -> str:
    value = _identifier(value)
    if value.casefold() in _SYSTEM_SCHEMAS or value.casefold().startswith("pg_"):
        raise DatabaseAccessError("system_catalog_forbidden")
    return value


def _host(value: str) -> tuple[str, bool]:
    try:
        address = ipaddress.ip_address(value)
        return value, address.is_loopback
    except ValueError:
        pass
    if not isinstance(value, str) or len(value) > 253 or not re.fullmatch(r"[A-Za-z0-9](?:[A-Za-z0-9.-]*[A-Za-z0-9])?", value):
        raise DatabaseAccessError("database_host_invalid")
    if any(not part or len(part) > 63 or part.startswith("-") or part.endswith("-") for part in value.split(".")):
        raise DatabaseAccessError("database_host_invalid")
    return value, value.casefold() == "localhost"


def _jdbc_options(dsn: str, spec: ConnectionSpec) -> dict[str, str]:
    """Translate the supported JDBC URL subset, never execute JDBC properties.

    Credentials may be supplied separately, as in JDBC's three-argument API.
    Ambiguous duplicate properties fail instead of silently changing targets or
    credentials. TLS, host validation and the bounded timeout are applied below.
    """
    if re.search(r"[\x00-\x20\x7f]", dsn) or re.search(r"%(?![0-9A-Fa-f]{2})", dsn):
        raise DatabaseAccessError("database_connection_invalid")
    url = urlsplit(dsn.removeprefix("jdbc:"))
    if not url.hostname or url.username is not None or url.password is not None or "#" in dsn:
        raise DatabaseAccessError("database_connection_invalid")
    if not url.path.startswith("/") or not url.path[1:] or "/" in url.path[1:]:
        raise DatabaseAccessError("database_connection_invalid")
    # url.port validates numeric ports; reject a dangling colon as well.
    if url.netloc.endswith(":"):
        raise DatabaseAccessError("database_connection_invalid")
    options = {"host": url.hostname, "port": str(url.port if url.port is not None else 5432),
               "dbname": unquote(url.path[1:], errors="strict")}
    allowed = {"user": "user", "password": "password", "sslmode": "sslmode",
               "connectTimeout": "connect_timeout", "connect_timeout": "connect_timeout"}
    for name, value in parse_qsl(url.query, keep_blank_values=True, strict_parsing=True, errors="strict"):
        key = allowed.get(name)
        if key is None:
            raise DatabaseAccessError("database_connection_option_forbidden")
        if key in options:
            raise DatabaseAccessError("database_connection_invalid")
        options[key] = value
    for key, value in (("user", spec.username),
                       ("password", spec.password.get_secret_value() if spec.password is not None else None)):
        if value is not None:
            if key in options and options[key] != value:
                raise DatabaseAccessError("database_connection_credentials_conflict")
            options[key] = value
    return options


def _resolved(spec: ConnectionSpec) -> dict:
    """Normalize endpoint options without echoing parser exceptions."""
    try:
        if spec.kind == "mysql_database":
            if spec.dsn is not None:
                dsn = spec.dsn.get_secret_value()
            else:
                host, _ = _host(spec.host)
                host = "[" + host + "]" if ":" in host else host
                dsn = ("mysql://" + quote(spec.username, safe="") + ":" +
                       quote(spec.password.get_secret_value() if spec.password else "", safe="") +
                       "@" + host + ":" + str(spec.port or 3306) + "/" + quote(spec.database, safe=""))
            # Reuse the adapter's strict driver flags and remote TLS validation.
            options = _connect_options(dsn, QueryLimits(timeout_ms=5000))
            try:
                host, _ = _host(options["host"])
                schema = _schema(spec.schema_name or options["database"])
                _identifier(options["database"])
                return {"dsn": dsn, "host": host, "port": options["port"],
                        "database": options["database"], "username": options["user"], "schema": schema}
            finally:
                options.clear()
        from psycopg.conninfo import conninfo_to_dict, make_conninfo
        if spec.dsn is not None:
            dsn = spec.dsn.get_secret_value().strip()
            options = (_jdbc_options(dsn, spec) if dsn.startswith(_GAUSS_JDBC_PREFIXES)
                       else conninfo_to_dict(dsn))
        else:
            options = {"host": spec.host, "port": str(spec.port or 5432), "dbname": spec.database,
                       "user": spec.username, "password": spec.password.get_secret_value() if spec.password else ""}
        try:
            # No service/passfile/options or multi-host DSNs: onboarding must not
            # execute startup SQL or resolve secrets from server-local files.
            if set(options) - {"host", "port", "dbname", "user", "password", "sslmode", "connect_timeout"}:
                raise DatabaseAccessError("database_connection_option_forbidden")
            if not options.get("host") or not options.get("user") or not options.get("dbname"):
                raise DatabaseAccessError("database_connection_fields_required")
            host, loopback = _host(options["host"])
            port = int(options.get("port") or 5432)
            if not 1 <= port <= 65535:
                raise DatabaseAccessError("database_connection_invalid")
            _identifier(options["dbname"])
            schema = _schema(spec.schema_name or "public")
            # An explicit form choice overrides the URL; omission preserves the
            # URL mode or the existing defaults. Never retry with weaker TLS.
            if options.get("sslmode") and options["sslmode"] not in _GAUSS_SSL_MODES:
                raise DatabaseAccessError("database_sslmode_invalid")
            mode = spec.sslmode or options.get("sslmode") or ("disable" if loopback else "verify-full")
            options.update(port=str(port), sslmode=mode, connect_timeout=str(CONNECTION_TIMEOUT_SECONDS))
            return {"dsn": make_conninfo(**options), "host": host, "port": port,
                    "database": options["dbname"], "username": options["user"], "schema": schema}
        finally:
            options.clear()
    except DatabaseAccessError:
        raise
    except Exception:
        raise DatabaseAccessError("database_connection_invalid") from None


def build_dsn(spec: ConnectionSpec) -> str:
    """Return an ephemeral DSN; callers must never persist or log this value."""
    return _resolved(spec)["dsn"]


def safe_connection_details(spec: ConnectionSpec) -> dict:
    resolved = _resolved(spec)
    resolved.pop("dsn")
    return {"kind": spec.kind, **resolved}


def normalize_selected_tables(spec: ConnectionSpec, names: list[str]) -> list[str]:
    """Bind manual/imported selections to the explicitly chosen schema."""
    schema = _resolved(spec)["schema"]
    if not isinstance(names, list) or not names or len(names) > TABLE_LIMIT:
        raise DatabaseAccessError("database_table_selection_invalid")
    result = []
    for name in names:
        if not isinstance(name, str):
            raise DatabaseAccessError("database_identifier_invalid")
        parts = name.split(".")
        if len(parts) == 1:
            parts = [schema, parts[0]]
        if len(parts) != 2 or _schema(parts[0]) != schema:
            raise DatabaseAccessError("database_table_schema_mismatch")
        qualified = parts[0] + "." + _identifier(parts[1])
        if qualified not in result:
            result.append(qualified)
    return result


@contextmanager
def _read_only_cursor(spec: ConnectionSpec):
    resolved = _resolved(spec)
    connection = cursor = None
    options = {}
    try:
        if spec.kind == "mysql_database":
            import pymysql
            options = _connect_options(resolved["dsn"], QueryLimits(timeout_ms=5000))
            connection = pymysql.connect(**options)
        else:
            import psycopg
            connection = psycopg.connect(resolved["dsn"], connect_timeout=CONNECTION_TIMEOUT_SECONDS)
        cursor = connection.cursor()
        if spec.kind == "mysql_database":
            cursor.execute("SET SESSION max_execution_time = 5000")
            cursor.execute("SET SESSION lock_wait_timeout = 1")
            cursor.execute("SET SESSION innodb_lock_wait_timeout = 1")
            cursor.execute("SET SESSION TRANSACTION READ ONLY")
            cursor.execute("START TRANSACTION READ ONLY")
            cursor.execute("SELECT @@session.transaction_read_only")
            read_only = cursor.fetchone()
        else:
            cursor.execute("BEGIN TRANSACTION READ ONLY")
            cursor.execute("SET LOCAL search_path TO pg_catalog")
            cursor.execute("SET LOCAL statement_timeout = '5000ms'")
            cursor.execute("SET LOCAL lock_timeout = '1000ms'")
            cursor.execute("SELECT current_setting('transaction_read_only')")
            read_only = cursor.fetchone()
        if not read_only or str(read_only[0]).lower() not in {"1", "on", "true"}:
            raise DatabaseAccessError("read_only_transaction_unverified")
        yield cursor, resolved["schema"]
    except DatabaseAccessError:
        raise
    except Exception:
        raise DatabaseAccessError("database_connection_failed") from None
    finally:
        failed = sys.exc_info()[0] is not None
        cleanup_failed = False
        if cursor is not None:
            try:
                cursor.close()
            except Exception:
                cleanup_failed = True
            finally:
                # PyMySQL keeps interpolated SQL and fetched rows on cursors.
                for attribute in ("_executed", "_rows"):
                    if hasattr(cursor, attribute):
                        setattr(cursor, attribute, None)
        if connection is not None:
            try:
                connection.rollback()
            except Exception:
                cleanup_failed = True
            finally:
                try:
                    connection.close()
                except Exception:
                    cleanup_failed = True
        options.clear()
        resolved.clear()
        if cleanup_failed and not failed:
            raise DatabaseAccessError("database_connection_close_failed") from None


def test_connection(spec: ConnectionSpec) -> dict:
    with _read_only_cursor(spec) as (cursor, _):
        cursor.execute("SELECT 1")
        row = cursor.fetchone()
        if not row or row[0] != 1:
            raise DatabaseAccessError("database_connection_probe_failed")
    return {"status": "ok"}


def discover_tables(spec: ConnectionSpec) -> dict:
    with _read_only_cursor(spec) as (cursor, schema):
        if spec.kind == "mysql_database":
            cursor.execute(
                "SELECT TABLE_SCHEMA, TABLE_NAME, TABLE_TYPE FROM information_schema.TABLES "
                "WHERE TABLE_SCHEMA = %s AND TABLE_TYPE IN ('BASE TABLE', 'VIEW') "
                "ORDER BY TABLE_NAME LIMIT %s", (schema, TABLE_LIMIT + 1))
        else:
            cursor.execute(
                "SELECT table_schema, table_name, table_type FROM information_schema.tables "
                "WHERE table_schema = %s AND table_type IN ('BASE TABLE', 'VIEW') "
                "ORDER BY table_name LIMIT %s", (schema, TABLE_LIMIT + 1))
        rows = cursor.fetchmany(TABLE_LIMIT + 1)
        tables = []
        for row in rows[:TABLE_LIMIT]:
            if len(row) != 3 or row[0] != schema or row[2] not in {"BASE TABLE", "VIEW"}:
                raise DatabaseAccessError("database_catalog_invalid")
            name = _identifier(row[1])
            tables.append({"schema": schema, "name": name, "qualified_name": schema + "." + name,
                           "table_type": row[2]})
        truncated = len(rows) > TABLE_LIMIT
    return {"tables": tables, "truncated": truncated}


def parse_table_list_excel(content: bytes) -> list[str]:
    """Read the first worksheet, never retain cells or execute Excel formulas.

    Recognized headers can put table/schema names in separate columns. Otherwise
    the first column is treated as the table list. Unqualified names are later
    resolved against the connection's selected schema.
    """
    if not isinstance(content, bytes) or not content or len(content) > MAX_EXCEL_BYTES:
        raise DatabaseAccessError("table_list_size_limit")
    workbook = None
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 128 or sum(entry.file_size for entry in entries) > MAX_EXCEL_EXPANDED_BYTES:
                raise DatabaseAccessError("table_list_size_limit")
            if any("vbaproject" in entry.filename.lower() or "externallinks/" in entry.filename.lower() for entry in entries):
                raise DatabaseAccessError("table_list_active_content_forbidden")
            for entry in entries:
                if entry.filename.lower().endswith((".xml", ".rels")):
                    with archive.open(entry) as stream:
                        tail = b""
                        while block := stream.read(65536):
                            # NUL stripping also catches UTF-16/UTF-32 entity
                            # declarations before the XML parser can expand them.
                            normalized = (tail + block).replace(b"\x00", b"").upper()
                            if b"<!DOCTYPE" in normalized or b"<!ENTITY" in normalized:
                                raise DatabaseAccessError("table_list_xml_entities_forbidden")
                            tail = (tail + block)[-64:]
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if not workbook.worksheets:
            raise DatabaseAccessError("table_list_empty")
        sheet = workbook.worksheets[0]
        if sheet.max_row and sheet.max_row > TABLE_LIMIT + 1 or sheet.max_column and sheet.max_column > 32:
            raise DatabaseAccessError("table_list_size_limit")
        table_headers = {"table", "table_name", "tablename", "qualified_name", "表", "表名", "数据表", "数据表名"}
        schema_headers = {"schema", "schema_name", "数据库", "模式"}
        result = []
        table_column = 0
        schema_column = None
        for index, row in enumerate(sheet.iter_rows(max_row=TABLE_LIMIT + 2, max_col=32)):
            if index > TABLE_LIMIT:
                if any(cell.value is not None for cell in row):
                    raise DatabaseAccessError("table_list_size_limit")
                break
            if index == 0:
                normalized = [str(cell.value or "").strip().lower() for cell in row]
                columns = [i for i, value in enumerate(normalized) if value in table_headers]
                if columns:
                    table_column = columns[0]
                    schemas = [i for i, value in enumerate(normalized) if value in schema_headers]
                    schema_column = schemas[0] if schemas else None
                    continue
            table_cell = row[table_column]
            schema_cell = row[schema_column] if schema_column is not None else None
            if table_cell.data_type == "f" or schema_cell is not None and schema_cell.data_type == "f":
                raise DatabaseAccessError("table_list_formula_forbidden")
            if table_cell.value is None or str(table_cell.value).strip() == "":
                continue
            if not isinstance(table_cell.value, str):
                raise DatabaseAccessError("database_identifier_invalid")
            value = table_cell.value.strip()
            parts = value.split(".")
            if schema_cell is not None and schema_cell.value is not None:
                schema = _schema(str(schema_cell.value).strip())
                if len(parts) == 1:
                    parts = [schema, parts[0]]
                elif parts[0] != schema:
                    raise DatabaseAccessError("database_table_schema_mismatch")
            if len(parts) not in {1, 2}:
                raise DatabaseAccessError("database_identifier_invalid")
            if len(parts) == 2:
                _schema(parts[0])
            value = ".".join(_identifier(part) for part in parts)
            if value not in result:
                result.append(value)
            if len(result) > TABLE_LIMIT:
                raise DatabaseAccessError("table_list_size_limit")
        if not result:
            raise DatabaseAccessError("table_list_empty")
        return result
    except DatabaseAccessError:
        raise
    except Exception:
        raise DatabaseAccessError("table_list_invalid") from None
    finally:
        if workbook is not None:
            workbook.close()
