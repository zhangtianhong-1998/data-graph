"""Read-only Excel input; no formula execution, network access, or staging files."""
from __future__ import annotations
from pathlib import Path
from datetime import date, datetime, time as daytime
import re
import sqlite3
import time
import zipfile
from xml.etree import ElementTree as ET
import openpyxl
from lightgraph.contracts import TableSpec, ColumnSpec, QueryLimits
from lightgraph.errors import LightGraphError
from lightgraph.query import QueryCompiler
from lightgraph.query.compiler import ident
from lightgraph.query.predicate import ScopeConjunction
from .memory import MemorySource
from .budget import BudgetedCursor

class ExcelSource(MemorySource):
    def __init__(self, path: str | Path, *, source_id: str = "excel", limits: QueryLimits | None = None, header_row: int = 1, include_hidden: bool = False, max_file_bytes: int = 200 * 1024 * 1024, max_rows: int = 200000, max_columns: int = 2000, max_memory_bytes: int = 128 * 1024 * 1024, query_budget=None):
        path = Path(path)
        self.issues: list[dict] = []
        self.issue_counts: dict[str, int] = {}
        self._pending_formula_columns: dict[str, set[str]] = {}
        if not path.is_file() or path.stat().st_size > max_file_bytes: raise LightGraphError("excel_file_limit")
        if path.suffix.lower() != ".xlsx": raise LightGraphError("excel_format_unsupported")
        if header_row < 1 or header_row > 100: raise LightGraphError("excel_header_invalid")
        try:
            with zipfile.ZipFile(path) as archive:
                members = archive.infolist()
                if sum(x.file_size for x in members) > max_memory_bytes * 4 or len(members) > 10000:
                    raise LightGraphError("excel_expansion_limit")
                if any("vbaproject" in x.filename.lower() or "externallinks/" in x.filename.lower() for x in members):
                    raise LightGraphError("excel_active_content_forbidden")
                # Reject XML entity declarations before XML readers see them.
                for member in members:
                    if member.filename.endswith(".xml"):
                        with archive.open(member) as stream:
                            tail = b""
                            while block := stream.read(65536):
                                data = tail + block
                                if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
                                    raise LightGraphError("excel_xml_entities_forbidden")
                                tail = data[-32:]
            book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
            values_book = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
        except LightGraphError: raise
        except Exception: raise LightGraphError("excel_parse_failed") from None
        tables, rows = [], {}
        try:
            for ws in book.worksheets:
                if ws.sheet_state != "visible" and not include_hidden: continue
                if (ws.max_column or 0) > max_columns: raise LightGraphError("excel_column_limit")
                # Count while reading too: dimensions in source files may lie.
                source_rows = ws.iter_rows(min_row=header_row)
                value_rows = values_book[ws.title].iter_rows(min_row=header_row)
                header = next(source_rows, ()); next(value_rows, ())
                if not header: continue
                if any(c.data_type == "f" for c in header): raise LightGraphError("excel_formula_header")
                last = max((i for i,c in enumerate(header) if c.value is not None), default=-1)
                if last < 0: continue
                names = [str(c.value) if c.value is not None else "" for c in header[:last+1]]
                if any(not n for n in names): raise LightGraphError("excel_merged_or_missing_header_requires_schema")
                if len(names) != len(set(names)): raise LightGraphError("excel_duplicate_header")
                table = TableSpec(source_id=source_id, schema="public", name=ws.title, columns=[ColumnSpec(name=n, comment="Uploaded Excel header") for n in names])
                hidden_rows, hidden_columns = self._hidden_positions(path, ws, include_hidden)
                kept = [i for i in range(len(names)) if i + 1 not in hidden_columns]
                table.columns = [table.columns[i] for i in kept]
                if not kept: continue
                tables.append(table)
                def iterate(sheet=ws, sr=source_rows, vr=value_rows, keep=kept, hidden=hidden_rows, column_names=tuple(names)):
                    for index, (raw, cached) in enumerate(zip(sr, vr), header_row + 1):
                        if index - header_row > max_rows: raise LightGraphError("excel_row_limit")
                        if index in hidden: continue
                        if len(raw) > max_columns: raise LightGraphError("excel_column_limit")
                        converted = []
                        for i in keep:
                            cell = raw[i] if i < len(raw) else None
                            cache = cached[i] if i < len(cached) else None
                            if cell is None: value = None
                            elif cell.data_type == "f":
                                value = cache.value if cache else None
                                if value is None:
                                    self._pending_formula_columns.setdefault(f"public.{sheet.title}", set()).add(column_names[i])
                                    self.issue_counts["formula_cache_missing"] = self.issue_counts.get("formula_cache_missing", 0) + 1
                                    if len(self.issues) < 32:
                                        self.issues.append({"code": "formula_cache_missing", "sheet": sheet.title, "row": index, "column": i + 1})
                            else:
                                value = cell.value
                                if value is None and cell.data_type == "inlineStr": value = ""
                            if isinstance(value, (date, datetime, daytime)): value = value.isoformat()
                            if isinstance(value, int) and not isinstance(value, bool) and cell is not None and re.fullmatch("0+", cell.number_format or ""):
                                value = str(value).zfill(len(cell.number_format))
                            converted.append(value)
                        if any(v is not None for v in converted): yield tuple(converted)
                rows[table.qualified_name] = iterate()
            if not tables: raise LightGraphError("excel_no_visible_tables")
            super().__init__(tables, rows, limits=limits, max_memory_bytes=max_memory_bytes, query_budget=query_budget)
            self._infer_loaded_types()
        except LightGraphError:
            if hasattr(self, "_db"): self.close()
            raise
        except Exception:
            if hasattr(self, "_db"): self.close()
            raise LightGraphError("excel_parse_failed") from None
        finally: book.close(); values_book.close()

    def _infer_loaded_types(self):
        """Classify all imported non-null cells using value-free SQL aggregates.

        SQLite's BLOB affinity preserves numeric cells versus textual codes. A
        mixed column remains projectable for discovery but cannot be filtered
        through typed predicates until its interpretation is resolved.
        """
        shapes = ("integer", "real", "text", "blob")
        for table in self._tables:
            for start in range(0, len(table.columns), 128):
                columns = table.columns[start:start + 128]
                projections = [f"MAX(CASE WHEN typeof({ident(column.name)}) = '{shape}' THEN 1 ELSE 0 END)"
                               for column in columns for shape in shapes]
                sql = f"SELECT {', '.join(projections)} FROM {ident(table.schema_name)}.{ident(table.name)}"
                cursor = BudgetedCursor(self._db.cursor(), self.query_budget, cancel=self._db.interrupt)
                began = time.monotonic()
                self._db.set_progress_handler(lambda: int(time.monotonic() > min(
                    self.query_budget.deadline, began + self.limits.timeout_ms / 1000)), 500)
                try:
                    cursor.execute(sql)
                    self.stats.queries += 1
                    counts = cursor.fetchone()
                    for i, column in enumerate(columns):
                        present = {shape for shape, count in zip(shapes, counts[i * 4:i * 4 + 4]) if count}
                        if not present or present == {"text"}: column.data_type = "text"
                        elif present == {"integer"}: column.data_type = "integer"
                        elif present <= {"integer", "real"}: column.data_type = "real"
                        else: column.data_type = "mixed"
                except sqlite3.Error:
                    self.query_budget.check_time()
                    raise LightGraphError("excel_type_inference_failed") from None
                finally:
                    cursor.close(); self._db.set_progress_handler(None, 0)
                    self.stats.elapsed_seconds += time.monotonic() - began
        self.compiler = QueryCompiler(self._tables, dialect="sqlite")

    def project(self, table, columns, predicate=None, distinct=True, limit=None):
        required=set(columns)
        def predicate_columns(value):
            if isinstance(value, ScopeConjunction):
                return predicate_columns(value.user_predicate) | predicate_columns(value.partition_filter)
            if not isinstance(value,dict): return set()
            logical = set(value) & {"all", "any", "and", "not"}
            if logical:
                result = set()
                for op in logical:
                    children = [value[op]] if op == "not" else value[op]
                    if isinstance(children, list):
                        for child in children: result.update(predicate_columns(child))
                return result
            if "column" in value: return {value["column"]}
            return set(value)
        required.update(predicate_columns(predicate))
        if required & self._pending_formula_columns.get(table.qualified_name,set()):
            # A missing cache is unknown, never a real NULL. Publishing its
            # incomplete projection could wrongly retract previously known facts.
            raise LightGraphError("formula_cache_missing_in_projection")
        yield from super().project(table,columns,predicate,distinct,limit)

    @staticmethod
    def _hidden_positions(path, ws, include_hidden):
        if include_hidden: return set(), set()
        hidden_rows, hidden_columns = set(), set()
        # read_only worksheets omit dimension metadata. Inspect XML without
        # materializing the sheet, using the internal path supplied by openpyxl.
        with zipfile.ZipFile(path) as archive:
            with archive.open(ws._worksheet_path) as stream:
                for event, elem in ET.iterparse(stream, events=("end",)):
                    tag = elem.tag.rsplit("}",1)[-1]
                    if tag == "row" and elem.attrib.get("hidden") in ("1", "true"):
                        hidden_rows.add(int(elem.attrib["r"]))
                    elif tag == "col" and elem.attrib.get("hidden") in ("1", "true"):
                        hidden_columns.update(range(int(elem.attrib["min"]), int(elem.attrib["max"]) + 1))
                    elem.clear()
        return hidden_rows, hidden_columns
    def capabilities(self):
        result = super().capabilities(); result.product = "excel_memory"
        return result
